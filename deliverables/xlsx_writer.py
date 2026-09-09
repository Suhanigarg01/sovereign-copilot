"""
Writes engineering-calculation-style spreadsheets with steps shown, not just
a final number — satisfies "calculations with steps shown".
"""
from openpyxl import Workbook
from openpyxl.styles import Font


def generate_calc_sheet(output_path: str, title: str, headers: list, rows: list, formula_note: str = None) -> str:
    """
    rows: list of lists, same column count as headers. Any cell starting with
    "=" is written as a live Excel formula (e.g. "=B2*C2").
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Calculation"

    ws.merge_cells("A1:%s1" % chr(ord("A") + len(headers) - 1))
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)

    header_row = 3
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = Font(bold=True)

    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    if formula_note:
        note_row = header_row + len(rows) + 2
        ws.cell(row=note_row, column=1, value=f"Note: {formula_note}").font = Font(italic=True, size=9)

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(length + 4, 40)

    wb.save(output_path)
    return output_path
