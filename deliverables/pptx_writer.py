"""Writes a simple summary deck from structured slide content."""
from pptx import Presentation
from pptx.util import Inches, Pt


def generate_summary_deck(output_path: str, title: str, slides: list) -> str:
    """
    slides: list of {"heading": str, "bullets": [str, ...]}
    """
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = "Generated on-premise by the AI Workbench"

    bullet_layout = prs.slide_layouts[1]
    for s in slides:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = s["heading"]
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(s["bullets"]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = bullet
            p.font.size = Pt(18)

    prs.save(output_path)
    return output_path
