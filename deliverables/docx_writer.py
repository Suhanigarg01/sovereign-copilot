"""
Turns agent findings into an actual approval-note Word document, not a chat
reply. This is the artifact judges will open at the end of the demo.
"""
from datetime import date
from docx import Document
from docx.shared import Pt


def generate_approval_note(
    output_path: str,
    subject: str,
    findings: list,
    recommendation: str,
    prepared_by: str = "AI Workbench (draft — pending human review)",
    reference_docs: list = None,
) -> str:
    doc = Document()

    title = doc.add_heading("APPROVAL NOTE", level=0)
    title.alignment = 1  # center

    doc.add_paragraph(f"Date: {date.today().isoformat()}")
    doc.add_paragraph(f"Subject: {subject}")

    doc.add_heading("Key Findings", level=1)
    for f in findings:
        doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(recommendation)

    if reference_docs:
        doc.add_heading("Referenced SOPs / Manuals", level=1)
        for r in reference_docs:
            doc.add_paragraph(r, style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph(f"Prepared by: {prepared_by}")
    footer.runs[0].font.size = Pt(9)
    footer2 = doc.add_paragraph("This note was drafted by an on-premise AI assistant and requires sign-off "
                                 "by an authorized reviewer before action is taken.")
    footer2.runs[0].font.size = Pt(9)
    footer2.runs[0].italic = True

    doc.save(output_path)
    return output_path
